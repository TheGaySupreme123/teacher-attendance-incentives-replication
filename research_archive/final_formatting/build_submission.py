"""Prepare a submission copy of Shay's edited paper; original sources stay unchanged."""
from pathlib import Path
from copy import deepcopy
import json,re
from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

HERE=Path(__file__).resolve().parent
OUT=HERE.parent
ROOT=OUT.parents[1]
SOURCE=HERE/'source_google_doc.docx'
APPENDIX=ROOT/'seminar_paper_drafts/01_Main_Paper_Academic_Revision_A4.docx'
CONFIG=json.loads((HERE/'submission_text.json').read_text())
doc=Document(SOURCE)
changes=[]

def font(run):
    run.font.name='Times New Roman';run.font.size=Pt(12);run.font.color.rgb=RGBColor(0,0,0)
    rpr=run._r.get_or_add_rPr();rf=rpr.find(qn('w:rFonts'))
    if rf is not None:
        for key in list(rf.attrib):
            if 'theme' in key.lower(): del rf.attrib[key]
        for a in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+a),'Times New Roman')

def style(name,bold=False):
    s=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    s.font.name='Times New Roman';s.font.size=Pt(12);s.font.bold=bold;s.font.color.rgb=RGBColor(0,0,0)
    pf=s.paragraph_format;pf.line_spacing=1.5;pf.space_before=Pt(0);pf.space_after=Pt(0)
    for b in list(s._element.iter(qn('w:pBdr'))):b.getparent().remove(b)
    return s

normal=style('Normal');title=style('Title',True);head=style('Heading 1',True)
head.paragraph_format.space_before=Pt(9);head.paragraph_format.space_after=Pt(3)
head.paragraph_format.keep_with_next=True
for sec in doc.sections:
    sec.page_width=Mm(210);sec.page_height=Mm(297)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    sec.header_distance=sec.footer_distance=Inches(.4)

# Remove empty paragraphs and manual pagination left by interactive editing.
for p in list(doc.paragraphs):
    for br in list(p._p.xpath('.//w:br[@w:type="page"]')):br.getparent().remove(br)
    if not p.text.strip() and not p._p.xpath('.//w:drawing|.//w:object|.//m:oMath'):
        p._p.getparent().remove(p._p)

for p in doc.paragraphs:
    pf=p.paragraph_format
    p.style=normal
    pf.line_spacing=1.5;pf.space_before=Pt(0);pf.space_after=Pt(0)
    pf.first_line_indent=Inches(.5);pf.left_indent=pf.right_indent=Inches(0)
    pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;pf.widow_control=True
    pf.keep_with_next=False;pf.keep_together=False;pf.page_break_before=False
    for r in p.runs:font(r)

for p in doc.paragraphs:
    if p.text.startswith('This paper replicates and extends selected findings'):
        old=p.text;p.text=CONFIG['abstract'];changes.append({'location':'Abstract','before':old,'after':p.text,'reason':'Under 150 words; removes repeated detail and avoids treating separate outcome p-values as a formal cross-outcome test.'})
        for r in p.runs:font(r)
        p.paragraph_format.first_line_indent=Inches(0)
    if p.text.startswith('The estimates below use the main sample'):
        # This paragraph defines the extension sample and belongs at its opening.
        sample=p

# Put the extension sample under its own section, retaining every word.
section3=next(p for p in doc.paragraphs if p.text=='3. Extension')
section3._p.addnext(sample._p)
changes.append({'location':'Extension sample paragraph','reason':'Moved unchanged from the end of replication to the start of Section 3.'})

refs=next(p for p in doc.paragraphs if p.text=='References')
aihead=refs.insert_paragraph_before('5. AI Use Statement')
aip=refs.insert_paragraph_before(CONFIG['ai_statement'])
for p in [aihead,aip]:
    p.style=normal;p.paragraph_format.line_spacing=1.5;p.paragraph_format.space_after=Pt(0)
    for r in p.runs:font(r)
changes.append({'location':'5. AI Use Statement','after':CONFIG['ai_statement'],'reason':'Required disclosure based on the retained project record.'})

in_refs=False
for i,p in enumerate(doc.paragraphs):
    t=p.text.strip();pf=p.paragraph_format
    if i==0:
        p.style=title;pf.first_line_indent=Inches(0);pf.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.keep_with_next=True
    elif i==1:
        pf.first_line_indent=Inches(0);pf.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.keep_with_next=True
    elif t=='Abstract':
        pf.first_line_indent=Inches(0);pf.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.keep_with_next=True
        for r in p.runs:r.bold=True
    elif re.match(r'^[1-5]\. ',t):
        p.style=head;pf.first_line_indent=Inches(0);pf.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.keep_with_next=True
        for r in p.runs:r.bold=True
    elif t=='References':
        in_refs=True;p.style=head;pf.page_break_before=True;pf.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.first_line_indent=Inches(0);pf.keep_with_next=True
        for r in p.runs:r.bold=True
    elif in_refs:
        pf.first_line_indent=Inches(-.5);pf.left_indent=Inches(.5);pf.alignment=WD_ALIGN_PARAGRAPH.LEFT;pf.keep_together=True
    elif re.match(r'^(Table [123]|Figure 1)$',t):
        pf.first_line_indent=Inches(0);pf.alignment=WD_ALIGN_PARAGRAPH.LEFT;pf.keep_with_next=True;pf.space_before=Pt(6)
    elif t.startswith(('Replication of Original','Extension of Original','Recorded Departure Effects')):
        pf.first_line_indent=Inches(0);pf.alignment=WD_ALIGN_PARAGRAPH.LEFT;pf.keep_with_next=True
    elif t.startswith('Note.'):
        pf.first_line_indent=Inches(0);pf.alignment=WD_ALIGN_PARAGRAPH.LEFT;pf.keep_together=True;pf.space_after=Pt(6)
    elif p._p.xpath('.//w:drawing'):
        pf.first_line_indent=Inches(0);pf.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.keep_with_next=True
    elif t=='I estimate':
        pf.first_line_indent=Inches(0);pf.keep_with_next=True

# Import only the actual appendices, not the earlier reference paper's main text.
source=Document(APPENDIX);active=False;appendix_blocks=[]
for el in source.element.body:
    if el.tag==qn('w:p'):
        p=Paragraph(el,source)
        if p.text.startswith('Appendix A Fixed sensitivity analyses'):active=True
        if active and p.text.strip():appendix_blocks.append(('p',p.text))
    elif active and el.tag==qn('w:tbl'):
        appendix_blocks.append(('table',[[c.text for c in r.cells] for r in Table(el,source).rows]))

in_d=False
for kind,value in appendix_blocks:
    if kind=='p':
        if value.startswith('Appendix D '):in_d=True
        if in_d and not value.startswith('Appendix D '):continue
        if value.startswith('Appendix '):
            p=doc.add_paragraph(value,style=head);p.paragraph_format.page_break_before=True
            p.paragraph_format.first_line_indent=Inches(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            if in_d:
                for text in CONFIG['appendix_d']:
                    q=doc.add_paragraph(text,style=normal);q.paragraph_format.line_spacing=1.5;q.paragraph_format.first_line_indent=Inches(.5);q.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
                    for r in q.runs:font(r)
                if CONFIG.get('repository_url'):
                    q=doc.add_paragraph('The paper and replication materials are available in the public ',style=normal)
                    q.paragraph_format.line_spacing=1.5;q.paragraph_format.first_line_indent=Inches(.5);q.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
                    for r in q.runs:font(r)
                    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),doc.part.relate_to(CONFIG['repository_url'],RT.HYPERLINK,is_external=True))
                    run=OxmlElement('w:r');props=OxmlElement('w:rPr');fonts=OxmlElement('w:rFonts')
                    for key in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+key),'Times New Roman')
                    props.append(fonts);size=OxmlElement('w:sz');size.set(qn('w:val'),'24');props.append(size)
                    color=OxmlElement('w:color');color.set(qn('w:val'),'000000');props.append(color)
                    underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');props.append(underline)
                    run.append(props);text=OxmlElement('w:t');text.text='GitHub repository';run.append(text);link.append(run);q._p.append(link)
                    font(q.add_run('.'))
        elif value.startswith('Table A'):
            number,caption=value.split(' ',2)[0:2],None
            m=re.match(r'^(Table A[12]) (.*)$',value)
            p=doc.add_paragraph(style=normal);p.add_run(m[1]).bold=True;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(6)
            q=doc.add_paragraph(style=normal);q.add_run(m[2]).italic=True;q.paragraph_format.keep_with_next=True
        else:
            value=value.replace('Notes. ','Note. ',1) if value.startswith('Notes. ') else value
            if value.startswith('The original teacher-score attendance'):
                value=value.replace('0.132 standard deviations','0.132 standard deviations per one-standard-deviation increase in baseline teacher score').replace('reduce it to 0.087,','reduce this interaction to 0.087,')
            p=doc.add_paragraph(value,style=normal);p.paragraph_format.line_spacing=1.5;p.paragraph_format.first_line_indent=Inches(.5);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            if value.startswith('Note.'):
                p.clear();p.add_run('Note.').italic=True;p.add_run(value[5:]);p.paragraph_format.first_line_indent=Inches(0);p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.keep_together=True;p.paragraph_format.space_after=Pt(6)
        for r in p.runs:font(r)
    else:
        t=doc.add_table(rows=0,cols=len(value[0]));t.autofit=False
        for row in value:
            cells=t.add_row().cells
            for cell,text in zip(cells,row):cell.text=text

# Consistent APA horizontal table rules and deliberate widths on A4.
width_pt=(210/25.4-2)*72
for ti,t in enumerate(doc.tables):
    t.autofit=False
    if ti==0:widths=[.265,.125,.205,.23,.175]
    elif ti==1:widths=[.36,.2133,.2133,.2134]
    elif ti==2:widths=[.25,.075,.215,.14,.12,.20]
    elif ti==3:widths=[.33,.16,.22,.13,.16]
    else:widths=[.48,.19,.33]
    for i,col in enumerate(t.columns):col.width=Pt(width_pt*widths[i])
    pr=t._tbl.tblPr
    tw=pr.find(qn('w:tblW'))
    if tw is not None:tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),str(round(width_pt*20)))
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();n=OxmlElement('w:cantSplit');trpr.append(n)
        if ri==0:
            n=OxmlElement('w:tblHeader');trpr.append(n)
        for ci,cell in enumerate(row.cells):
            cell.width=Pt(width_pt*widths[ci]);tcpr=cell._tc.get_or_add_tcPr()
            borders=tcpr.find(qn('w:tcBorders'))
            if borders is None:borders=OxmlElement('w:tcBorders');tcpr.append(borders)
            for child in list(borders):borders.remove(child)
            for edge in ['top','left','bottom','right','insideH','insideV']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'nil')
                if (edge=='top' and ri==0) or(edge=='bottom' and ri in [0,len(t.rows)-1]):el.set(qn('w:val'),'single');el.set(qn('w:sz'),'6');el.set(qn('w:color'),'000000')
                borders.append(el)
            margins=tcpr.find(qn('w:tcMar'))
            if margins is None:margins=OxmlElement('w:tcMar');tcpr.append(margins)
            for side in ['top','bottom','left','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'50');el.set(qn('w:type'),'dxa');margins.append(el)
            for p in cell.paragraphs:
                p.style=normal;pf=p.paragraph_format;pf.line_spacing=1.5;pf.space_after=pf.space_before=Pt(0);pf.first_line_indent=Inches(0);pf.left_indent=pf.right_indent=Inches(0)
                pf.keep_with_next=True;pf.keep_together=True
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:font(r);r.bold=ri==0

# Keep actual display width inside the printable area without stretching.
for shape in doc.inline_shapes:
    if shape.width>Pt(width_pt):
        ratio=Pt(width_pt)/shape.width;shape.width=Pt(width_pt);shape.height=round(shape.height*ratio)

for sec in doc.sections:
    for container in [sec.header,sec.footer]:
        for p in container.paragraphs:
            p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
            for r in p.runs:font(r)

doc.core_properties.title=doc.paragraphs[0].text
doc.core_properties.author=CONFIG['author']
doc.core_properties.subject='Track A replication and extension seminar paper'
doc.core_properties.comments=''
output=OUT/CONFIG['docx_name'];doc.save(output)
(HERE/'exact_submission_changes.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2))
print(output)
print('Abstract words:',len(CONFIG['abstract'].split()))

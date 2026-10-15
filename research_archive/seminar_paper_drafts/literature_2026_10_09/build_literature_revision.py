"""Build two A4 documents using the existing academic draft's design."""
from pathlib import Path
import sys,re,json
from docx.shared import Inches, Mm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE=Path(__file__).resolve().parent
BASE=HERE.parent
sys.path.insert(0,str(BASE))
import build_documents as old

original_base=old.base_doc
original_table=old.table
BODY_INCHES=210/25.4-2

def a4_base():
    doc=original_base()
    section=doc.sections[0]
    section.page_width=Mm(210); section.page_height=Mm(297)
    doc.styles['Title'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    return doc

def link(p,label,url):
    h=OxmlElement('w:hyperlink')
    h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    fonts=OxmlElement('w:rFonts')
    for a in ('ascii','hAnsi','cs','eastAsia'):fonts.set(qn('w:'+a),'Times New Roman')
    pr.append(fonts)
    color=OxmlElement('w:color');color.set(qn('w:val'),'000000');pr.append(color)
    u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u)
    r.append(pr); t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def runs(p,txt):
    pieces=re.split(r'(\*\*.*?\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^\)]+\)|https?://\S+)',txt)
    for value in pieces:
        if not value:continue
        if value.startswith('**'):p.add_run(value[2:-2]).bold=True
        elif value.startswith('*'):p.add_run(value[1:-1]).italic=True
        elif value.startswith('`'):p.add_run(value[1:-1])
        elif value.startswith('[') and re.match(r'^\[[^\]]+\]\(',value):
            label,url=re.match(r'^\[([^\]]+)\]\(([^\)]+)\)$',value).groups();link(p,label,url)
        elif value.startswith(('https://','http://')):
            target=value.rstrip('.,;');link(p,target,target);p.add_run(value[len(target):])
        else:p.add_run(value)

def fit_table(doc,heads,data,widths=None):
    widths=widths or [6.5/len(heads)]*len(heads)
    return original_table(doc,heads,data,[w*BODY_INCHES/6.5 for w in widths])

old.base_doc=a4_base;old.table=fit_table;old.text_runs=runs

def build(src,out,main):
    old.from_markdown(src,out,main=main)
    from docx import Document
    doc=Document(out)
    for shape in doc.inline_shapes:
        if shape.width>Inches(BODY_INCHES):
            ratio=Inches(BODY_INCHES)/shape.width
            shape.height=int(shape.height*ratio);shape.width=Inches(BODY_INCHES)
    if not main:
        doc.core_properties.subject='Annotated peer-reviewed literature guide for the seminar paper'
        for style in ['Normal','Heading 1','Heading 2','Heading 3','Title']:
            doc.styles[style].paragraph_format.line_spacing=1.15
        for index,p in enumerate(doc.paragraphs):
            if p.text.startswith('Citation.'):
                p.paragraph_format.keep_with_next=True
                p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after=Pt(5)
            if p.style.name.startswith('Heading'):
                p.paragraph_format.keep_with_next=True
            if index+1<len(doc.paragraphs) and len(p.text.split())<40 and doc.paragraphs[index+1].style.name=='Heading 2':
                p.paragraph_format.keep_with_next=True
    doc.save(out)
    return {'filename':out.name,'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'a4':True}

results=[
    build(HERE/'source/01_main_paper_literature_expanded.md',BASE/'01_Main_Paper_Literature_Expanded_A4.docx',True),
    build(HERE/'source/literature_guide_final.md',BASE/'05_Peer_Reviewed_Literature_Guide_A4.docx',False)
]
print(json.dumps(results))

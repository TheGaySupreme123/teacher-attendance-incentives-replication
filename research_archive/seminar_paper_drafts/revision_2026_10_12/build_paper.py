"""Build the Track A academic example from verified saved results.

Run with the bundled Python runtime. Empirical source files are read only.
"""
from pathlib import Path
import sys, re, json
from docx import Document
from docx.shared import Inches, Mm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = Path(__file__).resolve().parent
BASE = HERE.parent
sys.path.insert(0, str(BASE))
import build_documents as old

original_base = old.base_doc
original_table = old.table
original_exhibit = old.exhibit
original_equation = old.eq
WIDTH = 210 / 25.4 - 2
OUT = BASE / '01_Main_Paper_Academic_Revision_A4.docx'

def base():
    doc = original_base()
    s = doc.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    doc.styles['Title'].font.size = Pt(12)
    doc.styles['Title'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles['Title'].paragraph_format.space_after = Pt(8)
    normal = doc.styles['Normal'].paragraph_format
    normal.first_line_indent = Inches(.2)
    normal.space_after = Pt(4)
    normal.line_spacing = 1.5
    h1 = doc.styles['Heading 1'].paragraph_format
    h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.first_line_indent = Inches(0)
    h1.space_before, h1.space_after = Pt(10), Pt(5)
    h1.line_spacing = 1.5
    h2 = doc.styles['Heading 2']
    h2.font.bold, h2.font.italic = False, True
    h2.paragraph_format.first_line_indent = Inches(0)
    h2.paragraph_format.space_before, h2.paragraph_format.space_after = Pt(8), Pt(4)
    h2.paragraph_format.line_spacing = 1.5
    return doc

def link(p, label, url):
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run, prop = OxmlElement('w:r'), OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    for a in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
        fonts.set(qn('w:' + a), 'Times New Roman')
    prop.append(fonts)
    size = OxmlElement('w:sz'); size.set(qn('w:val'), '24'); prop.append(size)
    color = OxmlElement('w:color'); color.set(qn('w:val'), '000000'); prop.append(color)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); prop.append(underline)
    run.append(prop)
    t = OxmlElement('w:t'); t.text = label; run.append(t)
    h.append(run); p._p.append(h)

def runs(p, txt):
    for value in re.split(r'(\*\*.*?\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^\)]+\)|https?://\S+)', txt):
        if not value: continue
        if value.startswith('**'): p.add_run(value[2:-2]).bold = True
        elif value.startswith('*'): p.add_run(value[1:-1]).italic = True
        elif value.startswith('`'): p.add_run(value[1:-1])
        elif value.startswith('[') and re.match(r'^\[[^\]]+\]\(', value):
            label, url = re.match(r'^\[([^\]]+)\]\(([^\)]+)\)$', value).groups()
            link(p, label, url)
        elif value.startswith(('https://', 'http://')):
            target = value.rstrip('.,;'); link(p, target, target)
            p.add_run(value[len(target):])
        else: p.add_run(value)

def table(doc, heads, data, widths=None):
    widths = widths or [6.5 / len(heads)] * len(heads)
    t = original_table(doc, heads, data, [w * WIDTH / 6.5 for w in widths])
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent = Inches(0)
                p.paragraph_format.line_spacing = 1.15
                if heads[0] == 'Baseline group':
                    p.paragraph_format.keep_with_next = True
    return t

def caption(doc, txt):
    p = old.para(doc, txt, keep=True)
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(7), Pt(5)
    for r in p.runs: r.bold = True
    return p

def note(doc, txt):
    p = old.para(doc, 'Notes. ' + txt)
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(5), Pt(9)
    p.paragraph_format.keep_together = True
    return p

def equation(doc, key):
    original_equation(doc, key)
    if key == 'departure':
        m = doc.paragraphs[-1]._p.xpath('.//m:oMath')[0]
        m.append(old.mr('    (1)'))

def exhibit(doc, key):
    if key == 'replication':
        caption(doc, 'Table 1 Replication of original Table 2 Panel A and Table 10')
        data = [['School opening', '2\nPanel A', '0.21\n(0.03)', f"{old.ATT['estimate']:.3f}\n({old.ATT['se']:.3f})", '3,071\n113']]
        for z, label in zip(old.REPLICATION, ['Recorded departure', 'Government-school marker', 'Residual departure']):
            data.append([label, '10', f"{float(z['reference_csv_coef_rounded']):.2f}\n({float(z['reference_csv_se_rounded']):.2f})", f"{float(z['estimate']):.3f}\n({float(z['se_clustered']):.3f})", '2,197\n111'])
        table(doc, ['Outcome', 'Original\ntable', 'Published', 'Replicated', 'N\nSchools'], data, [1.95, .8, 1.2, 1.25, 1.3])
        note(doc, 'Original table numbers refer to Duflo, Hanna, and Ryan (2012). Coefficients are in probability units; school-clustered standard errors appear in parentheses. Attendance uses random visits in September 2003–February 2006; N counts visits. Departure uses children and the latest up to five observed open-school checks; N counts children. The government-school marker and residual retain the original reason-code convention. Models include an intercept and treatment assignment. Sources: original Tables 2 and 10, raw data, and Python reconstructions.')
    elif key == 'group_rates':
        caption(doc, 'Table 2 Extension of original Table 10 by baseline writing')
        data = []
        for g, label in [(0, 'Nonwriters'), (1, 'Writers')]:
            z = next(v for v in old.GROUP if v['outcome'] == 'drop' and int(v['pre_writ']) == g)
            cn, tn = int(z['n_control']), int(z['n_treated'])
            data.append([label, f"{old.pp(z['control_mean'])}%\n({round(float(z['control_mean']) * cn)}/{cn})", f"{old.pp(z['treated_mean'])}%\n({round(float(z['treated_mean']) * tn)}/{tn})", old.pp(z['treatment_effect']), old.interval(z['ci_low'], z['ci_high'])])
        table(doc, ['Baseline group', 'Control', 'Treatment', 'Difference\npp', '95% CI\npp'], data, [1.45, 1.22, 1.22, 1.05, 1.56])
        note(doc, 'Cell proportions are descriptive; parentheses give departures/children. Differences and confidence intervals are in percentage points. N = 2,195 unique children in 111 schools, excluding two conflicting writing records retained in Table 1. Writers took the written baseline test. Departure is zero attendance in the latest up to five observed open-school checks. Intervals use school-clustered covariance and 110 degrees of freedom. Treatment changes observation opportunities; these rates do not establish permanent educational exit. Source: reconstructed child cohort and interaction model.')
    elif key == 'main_effects':
        caption(doc, 'Figure 1 Departure effects and the baseline writing interaction')
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(); run.add_picture(str(old.ASSETS / 'main_effects.png'), width=Inches(WIDTH))
        for prop in run._r.xpath('.//wp:docPr'):
            prop.set('descr', 'Treatment effects on recorded departure: nonwriters 5.0 points, writers 26.1 points, writer-minus-nonwriter difference 21.1 points; all with school-clustered 95 percent confidence intervals.')
        note(doc, 'Dots show treatment effects and their direct difference in percentage points; lines show 95 percent confidence intervals. N = 2,195 children in 111 schools (57 treatment, 54 control). The model includes treatment, baseline writing, and their interaction. Intervals use school-clustered covariance with 110 degrees of freedom. The vertical line marks zero. Source: main interaction model.')
    elif key == 'outcome_family':
        caption(doc, 'Table 3 Extension of original Table 10 across departure outcomes')
        labels = {'drop': 'Recorded departure', 'gov': 'Government-school marker', 'drop2': 'Residual departure'}
        data = []
        for z in old.FAMILY:
            if z['specification'] == 'primary_consensus_no_fe':
                data.append([labels[z['outcome']], old.pp(z['estimate']), old.interval(z['ci_low'], z['ci_high']), old.pval(z['p_value']), old.pval(z['holm_p_value'])])
        table(doc, ['Outcome', 'Interaction\npp', '95% CI\npp', 'p', 'Holm p'], data, [2.1, 1.1, 1.7, .7, .9])
        note(doc, 'Interactions are writer-minus-nonwriter treatment effects in percentage points. N = 2,195 children in 111 schools. Intervals and raw p-values use school-clustered covariance with 110 degrees of freedom. Holm adjusts the three-outcome family. Wild-bootstrap p-values are .007 for total departure and .010 for residual departure; none is reported for the government marker. Destination measures are incomplete. Source: saved interaction results. Inference remains exploratory; no stars are used.')
    else:
        original_exhibit(doc, key)

old.base_doc, old.table, old.text_runs = base, table, runs
old.caption, old.note, old.eq, old.exhibit = caption, note, equation, exhibit

old.from_markdown(HERE / 'source/01_main_paper_academic_revision.md', OUT, main=True)
doc = Document(OUT)
in_refs, in_appendix = False, False
for index, p in enumerate(doc.paragraphs):
    if p.style.name == 'Title':
        p.text = 'Attendance Incentives and Recorded School Departure\nby Baseline Writing Status'
    if p.style.name == 'Heading 2' and re.match(r'^\d+ \d+ ', p.text):
        # Conventional subsection numbering follows the scholarly reference.
        p.text = re.sub(r'^(\d+) (\d+) ', r'\1.\2 ', p.text)
    if p.text == '[Author name]':
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Inches(0)
    if p.text == 'Abstract':
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        nextp = doc.paragraphs[index + 1]
        nextp.paragraph_format.first_line_indent = Inches(0)
    if p.text == 'References': in_refs = True
    if p.text.startswith('Appendix A'): in_refs = False; in_appendix = True
    if in_refs and p.style.name == 'Normal' and p.text:
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Inches(.25)
        p.paragraph_format.first_line_indent = Inches(-.25)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_together = True
    if in_appendix and p.style.name == 'Normal' and 'scripts/' in p.text:
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if p.text.startswith('The main statistical outputs are'):
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
doc.core_properties.subject = 'Track A replication and baseline writing extension of Duflo Hanna Ryan 2012'
doc.core_properties.comments = 'Academic reference example prepared from verified saved analyses'
doc.save(OUT)
print(json.dumps({'output': str(OUT), 'paragraphs': len(doc.paragraphs), 'tables': len(doc.tables), 'a4': True}))

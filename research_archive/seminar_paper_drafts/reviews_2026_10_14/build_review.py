from pathlib import Path
from collections import Counter
import hashlib
import json
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/Users/shayb/Downloads/112523-V1')
OUT = ROOT / 'seminar_paper_drafts/reviews_2026_10_14'
SOURCE = Path('/Users/shayb/.codex/attachments/53e7c588-1bdb-4295-8899-1e9d4bb647f0/Pasted text.txt')
TARGET = OUT / '06_Shay_Rewrite_Annotated_Review_A4.docx'
PROTECTED = [SOURCE,
 ROOT / 'seminar_paper_drafts/01_Main_Paper_Academic_Revision_A4.docx',
 ROOT / 'seminar_paper_drafts/01_Main_Paper_Academic_Revision_A4.pdf',
 ROOT / 'seminar_paper_drafts/revision_2026_10_12/source/01_main_paper_academic_revision.md']
before = {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in PROTECTED}
original = SOURCE.read_text()
lines = original.splitlines()
assert len(lines) == 53
NOTES = []

def note(line, category, anchor, text, equation=False):
    assert anchor in lines[line-1], (line, anchor)
    NOTES.append(dict(line=line, category=category, anchor=anchor, text=text, equation=equation))

note(5, 'Correction', 'randomize schools to monitoring via photographs',
 'Replace with “randomize schools to a package of camera monitoring and attendance-linked pay.” Assignment identifies the effect of the package; it does not identify a separate effect of photographs. Apply the same qualification to “monitoring caused” later in this excerpt. [E1: printed pp. 1242, 1246]')
note(5, 'Clarity', 'study whether monitoring affects teacher absenteeism',
 'For the opening account of the original paper, state the broader question about incentives, teacher absence, and learning; then give its headline findings (about 21 percentage points lower absence and 0.17 SD higher test scores), before identifying the attendance and departure results you replicate. This makes the original contribution and your selected replication scope clear. [E1: printed pp. 1241–1242]')
note(6, 'Clarity', 'the control group continued to receive surprise visits from researchers',
 'Replace with “Researchers made unannounced attendance visits in both treatment and control schools. Only treatment schools used the daily camera records to calculate attendance-linked pay; control teachers received fixed pay.” The existing sentence can imply that these research visits were unique to control schools. [E1: printed pp. 1246–1248; E2]')
note(7, 'Correction', 'I replicate the paper’s Table 2',
 'Specify “the pooled all-teacher result in Table 2, Panel A, and all three outcome rows of Table 10.” Our attendance reconstruction reproduces that selected pooled result, not every panel and time-period specification in original Table 2. Use this scope consistently in lines 8–9. [E1: printed pp. 1249–1250, 1274; E3]')
note(7, 'Requirement', 'in Python',
 'State the Python version and analysis-package versions, or point explicitly to an appendix/environment file that supplies them, with the reproduction commands. Track A asks for exact software, sample, years, variable definitions, and missing-data treatment; submitted tables and figures must be reproducible from code and a README. A public folder link is encouraged, rather than mandatory. This excerpt alone does not show whether those details already appear elsewhere. [E5: pp. 3–4]')
note(7, 'Clarity', 'continue with the extension on my own',
 'If you mean that the extension is your addition to the original authors’ analysis, say that directly. Avoid implying unassisted coding or writing: the project used AI assistance. The full paper’s required AI Use Statement should accurately identify the tools, model/version, dates, and roles; keep the transcripts. Its absence from this two-section excerpt is not itself a finding. [E5: pp. 2–3]')
note(9, 'Correction', 'As in the original table, schools were open in 78.7%',
 'Say “My reconstruction gives 78.7% in treatment and 58.0% in control; these round to the published 0.79 and 0.58.” Replace “I match all results” with “I match the selected Table 2, Panel A estimates at their reported precision.” The reconstructed decimals are more precise than those displayed in the article. [E1: Table 2; E3]')
note(9, 'Requirement', 'The 95% confidence interval is [14.9, 26.6]',
 'Add “percentage points.” The coefficient 0.207 and SE 0.030 are in probability units; the corresponding effect and SE are 20.7 and 3.0 percentage points. Use one explicit scale throughout the prose and each exhibit. Also change “SEs … in parentheses” if the final display is “SE 0.030” rather than “(0.030).” [E3; E5: p. 2]')
note(10, 'Correction', 'select the last five times they were checked when the school was open',
 'Say “the latest up to five observed checks when the school was open.” Some children have fewer than five. Zero summed recoded attendance defines recorded departure under the distributed rule; it does not establish permanent exit from education. Replace “then this child has left” with “then the child is classified as having departed under this rule.” [E2: reconstruction, lines 143–159; E4]')
note(10, 'Correction', 'The dependent variables are (1)',
 'List all three replicated outcomes: recorded departure; recorded government-school marker; and residual departure, equal to departure minus the marker. These correspond to the original Table 10 rows “Child left NFE,” “Child enrolled in government school,” and “Child dropped out.” Retain the safer residual label because the third row does not verify dropout from education. [E1: Table 10; E2; E3]')
note(11, 'Correction', 'merged on to random checks using these identifiers',
 'Correct the joins: roster rows merge to school random-check records on school ID and visit year/month/day; child enrollment and baseline writing status merge on school ID and child ID. The fitted reconstruction does not merge the school random-check file on child ID or impose “random check time period = baseline” on the writing-status join. Name the two roster inputs (roster_rema2012.dta and rosterfin.dta) and move detailed join rules to Appendix B if needed. [E2: reconstruction, lines 100–136]')
note(11, 'Correction', 'I code these as having left but the destination unknown',
 'Replace with “Attendance determines recorded departure first. A blank reason contributes no government-school marker; another selected record can still supply reason equal to ‘2’ for that child.” Blank reasons do not themselves create departure. The child-level marker requires recorded departure and at least one selected record with reason equal to “2”; it is not independently verified follow-up enrollment. Apply that definition to “we know” in this paragraph and later passages. [E2: reconstruction, lines 143–159]')
note(11, 'Requirement', 'Although parts of this process are unclear to me',
 'Replace this general uncertainty with the specific documented issues: undocumented attendance codes and the numeric/text reason-field append convention. Track A requires an exact account of variable construction and missing-data treatment. It is appropriate to say that codebook meanings remain unresolved, while being precise about the implemented recodes. [E2; E4; E5: p. 3]')
note(12, 'Correction', 'Treatment causes a statistically significant 8.4pp increase in departure',
 'Use “Assignment to the bundled program increased recorded departure by 8.4 percentage points (clustered SE 4.1 points; p = .044) and the recorded government-school marker by 10.1 points (SE 3.5 points).” The nominal 5% significance claim for total recorded departure is supported. Qualify the endpoints rather than implying verified educational withdrawal or transfer. [E3: Table 10 reproduction]')
note(13, 'Requirement', 'a −1.8 point change in residual departure',
 'Use “−1.8 percentage points (clustered SE 3.1 points; p = .564),” or ensure the referenced main table gives that uncertainty. The aggregate residual effect is not statistically distinguishable from zero. “All child-level reference values” should specify the reconstructed departure indicators, rather than every variable in the package. [E2; E3; E5: p. 2]')
note(16, 'Correction', 'examines whether writers and nonwriters differ in departures',
 'Specify “whether the treatment effect on recorded departure differs between baseline writers and nonwriters.” This is a difference between two treatment effects, rather than simply a difference in departure levels. Keep the hypothesis about readiness tentative: written-test status is a coarse baseline preparation proxy, not a validated readiness or permanent-ability measure. Replace “drop-out from education” with “departure without the recorded government-school marker,” because educational dropout is not observed. [E1; E2; E3]')
note(17, 'Requirement', 'In Table 4',
 'Choose final exhibit numbers and use them consistently. Here the model is assigned to Table 4, the main effects later to Table 2, and a placeholder to Table 3. Distinguish “original Table 10” from your own tables. A workable arrangement is Table 1 for replication, Table 2 for group treatment effects, and Table 3 for the three interaction outcomes; another consistent arrangement is acceptable. [E5: pp. 2–3]')
note(18, 'Correction', '(1)',
 'Replace the displayed equation with the model below. The fitted primary model has a common intercept, treatment, a standalone baseline-writing indicator, and their interaction. It has no school, child, or class fixed effects and no observation-specific treatment coefficients. School fixed effects would absorb school-level treatment assignment. Block/geographic indicators belong only to the separate robustness specification. [E2: model; E3: primary_consensus_no_fe]', equation=True)
note(19, 'Correction', 'departure rate among nonwriters in control schools (α_s)',
 'Following the corrected equation, the four cell means are α, α + τ, α + β, and α + β + τ + δ. Replace γ with β and remove coefficient subscripts. Your interpretation of τ, τ + δ, and δ is correct. For the writer effect, the SE uses Var(τ) + Var(δ) + 2Cov(τ, δ); δ alone has its own clustered SE. [E2; E3]')
note(20, 'Correction', 'known to be attending a government school at follow-up',
 'Replace with “classified as departed and carrying a government-school reason marker in at least one selected record.” The marker is an operational reason-field outcome, not an independently verified follow-up attendance measure. Residual departure is the remaining recorded departures, including missing or unrecorded destinations. [E2: reconstruction, lines 143–159]')
note(20, 'Correction', 'follow-up-up 9,999 Rademacher randomization draws',
 'Replace with “9,999 null-imposed school-level Rademacher wild-bootstrap-t draws, seed 20260922.” These are random signs applied at the school level to restricted residuals, not new random assignments of treatment. The tested null is δ = 0, while the treatment and writing main effects remain unrestricted. Delete the duplicated “follow-up-up.” [E2: scope, lines 200–243; E3: bootstrap summaries]')
note(20, 'Correction', 'p-values for τ + δ',
 'Replace with “p-values for the interaction δ, the writer-minus-nonwriter difference in treatment effects.” τ + δ is the writer treatment effect. Holm adjusts the three conventional school-clustered interaction p-values; the supplementary wild-bootstrap p-values are reported separately and were not Holm-adjusted. Preserve the caveat that the three-outcome adjustment does not account for the earlier search across extension directions. [E3]')
note(38, 'Correction', '.010',
 'The displayed table has two rows even though line 20 says three outcomes. Add the total recorded-departure interaction row: 21.1 percentage points; 95% CI [6.9, 35.3]; clustered p = .004; Holm p = .012; wild-bootstrap p = .007. Fill the government-marker Holm cell with .534 and its wild-bootstrap cell with “—”, defined as not estimated/reported. Its wild p-value must not be invented. This repairs the stated three-outcome display; alternatively, clearly cross-reference another exhibit that shows the total interaction. [E3]')
note(39, 'Correction', 'sum to the estimate for Departure in Table 2',
 'Specify “the government-marker and residual interaction point estimates sum to the total recorded-departure interaction: 4.8 + 16.3 = 21.1 percentage points.” They do not sum to the aggregate 8.4-point replication effect. The decomposition applies to point estimates on the same sample/design; SEs, confidence intervals, and p-values do not add. Use the final total-interaction table reference. [E2; E3]')
note(39, 'Requirement', 'Notes: Table reports interaction coefficient',
 'Add that confidence intervals and conventional p-values use SEs clustered over 111 schools, with 110 degrees of freedom; Holm covers the three δ tests; wild-bootstrap values use the stated null-imposed school procedure and are separate. Define all three outcomes and the dash. Explain any significance stars if you use them. Label estimates and confidence intervals as percentage points and use a publication-quality table in the final paper. The existing N = 2,195 and school count are correct. [E3; E5: p. 2]')
note(40, 'Correction', 'and Duflo, Hanna and Ryan (2012) report it as exploratory',
 'Delete or complete this fragment without attributing our exploratory departure interaction to the original authors. They did not report this interaction; the exploratory label applies to our extension after examining several directions. Replace “high achievers” with “baseline writers.” Residual departure really is the residual operational category; its limitation is that it does not identify dropout or any particular destination. [E1: Tables 7, 9, 10; E2; E3]')
note(41, 'Requirement', 'Table here: Table 3 about:',
 'Replace this placeholder with the numbered, titled table actually discussed. Likewise replace the Figure 1 placeholder on the next line with a numbered, titled figure using labeled units and uncertainty, or remove its references if it is not included. Table and figure numbering are separate sequences. A coefficient plot is recommended by the assignment and would suit these effects, but that specific figure type is not mandatory. [E5: pp. 2–3]')
note(43, 'Correction', '8.4 percentage points, as did the original study',
 'Say “My reconstruction gives 8.4 percentage points, which rounds to the published 0.08.” Keep “recorded departure” and “bundled program” here. The journal table displays a rounded 8-point effect; 8.4 is our reconstruction at finer precision. [E1: Table 10; E3]')
note(43, 'Requirement', '(Duflo, Hanna and Ryan (2012) table 10 (E112523V1))',
 'Use one author-date convention throughout. For example, “Duflo, Hanna, and Ryan (2012, Table 10)” in text avoids nested citations. Check the similar forms in lines 5, 44, and 49. Consistent APA or Chicago author-date is an explicit requirement. [E5: p. 1]')
note(43, 'Clarity', 'E112523V1',
 'For clear provenance, cite the replication dataset separately as Version 1, DOI 10.3886/E112523V1, rather than treating the package ID as the journal-table reference. This is a provenance suggestion; the assignment does not prescribe this exact dataset-entry format. [E6]')
note(44, 'Correction', 'based on their performance in a baseline test',
 'Use “by baseline written-versus-oral test status.” This extension uses the existing protocol indicator, not the child’s achieved score or a passing-score cutoff. Apply that wording to “performance,” “ability,” and “children who could write” later in the excerpt. Your statement that original Table 10 does not report this particular heterogeneity analysis is supported. [E1: Tables 7, 9, 10; E2; E3]')
note(46, 'Correction', 'leave out the two children in my sample of 2,195',
 'Replace with “The aggregate replication uses 2,197 children. I exclude the two eligible children with conflicting baseline-writing records before the interaction analysis, leaving 2,195.” Line 14 already states this correctly. The present wording could imply a further reduction to 2,193. [E2; E3]')
note(46, 'Correction', 'the standard errors would not be comparable',
 'Delete this reason. All SEs can be reported on the same probability/percentage-point scale with the same school clustering. The estimates differ in estimand and slightly in sample, so label those differences explicitly. Also reconcile “Column 1 replicates” with “I do not report … together.” Columns giving nonwriter and writer effects contain τ and τ + δ, not two interaction coefficients. [E2; E3]')
note(47, 'Requirement', '26.1 percentage point treatment-control difference',
 'These point estimates are correct. Supply their uncertainty in the referenced table or here: writer effect 26.1, 95% CI [10.0, 42.3]; nonwriter effect 5.0, 95% CI [−2.4, 12.4], all in percentage points. The direct difference 21.1, CI [6.9, 35.3], and its p-values are also correct. Do not infer heterogeneity merely from one subgroup being significant and the other not; the δ test already supplies the right comparison. [E3; E5: p. 2]')
note(47, 'Correction', 'as in the original paper',
 'Delete this attribution after the wild-bootstrap statement. School-clustered SEs reproduce the original inference convention; the wild-bootstrap and Holm checks are additions in this extension. [E1: Tables 2, 9, 10; E3]')
note(48, 'Correction', 'interaction with baseline ability is much larger',
 'Say “The residual-departure interaction has a larger point estimate than the government-marker interaction (16.3 versus 4.8 percentage points).” No direct residual-minus-government contrast test was reported, so this is a point-estimate comparison, not an established difference between the two effects. Use baseline writing status rather than ability. [E3]')
note(49, 'Correction', 'significantly higher gains in learning',
 'Replace with “larger estimated learning gains among baseline writers.” Original Table 9 reports subgroup estimates; it does not report a formal equality test establishing that writer gains exceed nonwriter gains. Also qualify Table 7 as a comparison among children who stayed, a sample selected after treatment. Its 14-versus-10-point comparison does not by itself explain why particular children departed. [E1: printed pp. 1271–1274]')
note(51, 'Correction', 'the first two sum to departure as coded originally',
 'Name the components explicitly: “The government-marker and residual-departure interaction estimates sum to the total recorded-departure interaction.” This avoids implying that total departure plus the government marker sums to departure. “Including missing” means departures with missing destination reasons, not that every missing attendance or reason is automatically a departure. [E2; E3]')
note(51, 'Correction', 'such a pattern is not able to explain the difference between the groups',
 'Replace with “The recorded government-school marker does not establish a larger treatment effect for writers than nonwriters; its interaction is 4.8 percentage points, with CI [−10.5, 20.2]. Missing destination records prevent a stronger conclusion about actual transfers.” A nonsignificant marker interaction does not rule out a transfer explanation, nor test whether the absolute treatment effect among writers is zero. [E3; E4]')
note(51, 'Clarity', 'I report all three to be clear about the limitation',
 'Frame these outcomes as a decomposition diagnostic within the single baseline-writing extension. Track A asks for one main point and one sensible change; this framing preserves the same heterogeneity question rather than opening a separate destination study. Your caveat about interpretation is useful. [E5: p. 3]')
note(52, 'Correction', 'some have a coded reason for departure that does not appear in the available codebooks',
 'Replace with “113 recorded-departure classifications involve attendance values whose meanings are not established by the supplied codebook.” These are undocumented attendance sentinel codes, not unknown departure-reason codes. The counts 145, 113, overlap 28, and union 230 are correct for the 2,195-child extension sample. Also correct “excluding children with departure coded” to “excluding the 113 ambiguous attendance-based departure classifications.” [E4]')
note(52, 'Correction', 'to minimize the observations with fewer than five checks',
 'Replace with “to test dependence on the final January–February 2006 endpoint records, whose appended reason values become blank under the distributed Stata convention.” The fixed December 31, 2005 cutoff was not chosen to minimize short histories; an earlier cutoff can reduce observed checks. Describe the regional check as adding existing school-block/geographic indicators, without calling them verified randomization strata. Keep the detailed results in the appendix. [E2: scope, lines 11–15, 38–41, 88–96; E4; E5: p. 3]')
note(52, 'Clarity', 'bounding relies on other observations in the sample',
 'Be precise: these bounds vary the ambiguous children’s binary departure classifications while holding the other observed records fixed. They are algebraic observed-sample ranges, not confidence intervals or proof of statistical significance under every coding. The exclusion check changes the sample. Retain the present warning that none repairs treatment-dependent observation windows. [E4]')
note(53, 'Correction', 'schools run by nonwriters',
 'Writer/nonwriter describes children, not teachers, principals, or the people running schools. Replace the four affected spans here: “children … who went to schools run by nonwriters” → “baseline nonwriter children”; “schools run by writers” → “baseline writer children”; “whether its principal was a writer” → “the child’s baseline writing status”; and the final “children who went to schools run by writers than nonwriters” → “baseline writers than baseline nonwriters.” The open-check means 15.4/20.4 and 15.2/21.9 are correctly reported for those child groups. [E2; E3; E4]')
note(53, 'Correction', 'reason for departure is missing for 66.4% of observations',
 'Specify the denominator: “Reason fields are blank in 7,112 of 10,707 selected open-school check records (66.4%).” This is a check-record share, not the share of the 2,195 children or the share of departures with unknown destinations. A blank reason on a nondeparture check is not necessarily missing destination information. [E4: measurement summary]')
note(53, 'Clarity', 'common treatment-proof measure of exiting education',
 'Use “a treatment-independent measure of permanent educational exit.” Your central warning is sound: treatment affects opportunities to observe children in open schools, so a causal effect on the recorded endpoint need not be a causal effect on leaving education. Keep the conclusion limited to recorded NFE departure and the child-level subgroup difference. [E2; E4]')

OUT.mkdir(parents=True, exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21)
sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
sec.left_margin, sec.right_margin = Cm(2.2), Cm(2.2)
sec.header_distance, sec.footer_distance = Cm(.9), Cm(.9)
normal = doc.styles['Normal']
normal.font.name, normal.font.size = 'Times New Roman', Pt(12)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(6)
for s in ['Title','Heading 1','Heading 2']:
    doc.styles[s].font.name = 'Times New Roman'
    doc.styles[s].font.color.rgb = RGBColor.from_string('1B1B1B')
doc.styles['Title'].font.size = Pt(19)
doc.styles['Heading 1'].font.size = Pt(15)
doc.styles['Heading 2'].font.size = Pt(13)
review = doc.styles.add_style('Review note', 1)
review.font.name, review.font.size = 'Times New Roman', Pt(11)
review.font.color.rgb = RGBColor.from_string('9C2424')
review.paragraph_format.line_spacing = 1.15
review.paragraph_format.space_before = Pt(3)
review.paragraph_format.space_after = Pt(8)
review.paragraph_format.left_indent = Cm(.25)
review.paragraph_format.right_indent = Cm(.15)
review.paragraph_format.keep_together = True
source_style = doc.styles.add_style('Original source', 1)
source_style.base_style = normal
source_style.font.color.rgb = RGBColor(0,0,0)
compact = doc.styles.add_style('Original table line', 1)
compact.base_style = source_style
compact.paragraph_format.line_spacing = 1.05
compact.paragraph_format.space_after = Pt(0)
compact.paragraph_format.keep_with_next = True

header = sec.header.paragraphs[0]
header.text = 'SHAY’S REWRITE • REVIEW COPY • 14 OCTOBER 2026'
header.runs[0].font.name = 'Times New Roman'
header.runs[0].font.size = Pt(9)
header.runs[0].font.color.rgb = RGBColor.from_string('666666')
footer = sec.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('Review page ').font.size = Pt(9)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.core_properties.title = 'Shay’s replication and extension — annotated review'
doc.core_properties.subject = 'Read-only source review against verified study and Track A requirements'
doc.core_properties.author = 'Codex — review annotations'

doc.add_heading('Review of your replication and extension', 0)
doc.add_paragraph('Your original wording appears in black. Red parenthetical notes identify the exact quoted span and suggest what you could change. No correction has been applied to your prose or existing paper. The original equation is rendered as editable Word math with its original terms; its suggested replacement appears separately. The pasted table is laid out in columns, retaining its original labels, values, and blank cells.')
doc.add_paragraph('Correction = a mismatch with the original article, implemented analysis, or recorded results. Requirement = an explicit Track A or final-paper rule. Clarity = an optional improvement or a useful qualification. Source-line numbers refer to the pasted text, not manuscript page numbers.')
doc.add_paragraph('Priority: correct the model and inference paragraph, the attendance/destination construction, and the child-level writing labels. Then finish the exhibits and cross-references. Most headline coefficients and sample counts agree with the saved analysis.')
doc.add_paragraph('The requirements are used as review criteria; they do not authorize rewriting or submitting your work. This is a review copy, so its length and formatting are not a proposed submission layout.')
doc.add_heading('Annotated original excerpt', 1)

def mr(text):
    r = OxmlElement('m:r')
    t = OxmlElement('m:t'); t.text = text; r.append(t)
    return r

def sub(base, index):
    node = OxmlElement('m:sSub')
    e = OxmlElement('m:e'); e.append(mr(base)); node.append(e)
    s = OxmlElement('m:sub'); s.append(mr(index)); node.append(s)
    return node

def equation(p, corrected=False):
    math = OxmlElement('m:oMath')
    if corrected:
        expr = [sub('Y','is'), mr(' = α + τ '), sub('T','s'), mr(' + β '),
                sub('W','is'), mr(' + δ ('), sub('T','s'), mr(' × '),
                sub('W','is'), mr(') + '), sub('ε','is')]
    else:
        expr = [sub('Y','isc'), mr(' = '), sub('α','s'), mr(' + '),
                sub('γ','sc'), mr(' + '), sub('λ','c'), mr(' + '),
                sub('τ','isc'), mr(' × '), sub('T','sc'), mr(' + '),
                sub('δ','isc'), mr(' × ('), sub('W','ic'), mr(' × '),
                sub('T','sc'), mr(') + '), sub('ε','ics'), mr('     (1)')]
    for x in expr: math.append(x)
    if corrected:
        for r in math.iter(qn('m:r')):
            rp = OxmlElement('w:rPr')
            color = OxmlElement('w:color'); color.set(qn('w:val'),'9C2424')
            rp.append(color); r.insert(0,rp)
    p._p.append(math)

byline = {}
for i,n in enumerate(NOTES,1):
    n['id'] = f'N{i:02d}'
    byline.setdefault(n['line'], []).append(n)
source_paragraphs = []
def add_notes(lineno):
    for n in byline.get(lineno,[]):
        p = doc.add_paragraph(style='Review note')
        p.add_run(f'({n["id"]} · {n["category"]} · source line {lineno} · “{n["anchor"]}”: ').bold = True
        p.add_run(n['text'])
        if n['equation']:
            p.add_run(' Suggested equation: ')
            equation(p, corrected=True)
        p.add_run(')')

for lineno,text in enumerate(lines,1):
    if 21 <= lineno <= 38:
        if lineno != 21:
            continue
        table = doc.add_table(rows=3, cols=6)
        table.autofit = False
        widths = [3.7, .9, 3.4, 1.6, 2.5, 4.5]
        for rowidx,row in enumerate(table.rows):
            trpr = row._tr.get_or_add_trPr()
            nobreak = OxmlElement('w:cantSplit'); trpr.append(nobreak)
            if rowidx == 0:
                repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
            for colidx,cell in enumerate(row.cells):
                cell.width = Cm(widths[colidx])
                source_line = 21 + rowidx * 6 + colidx
                p = cell.paragraphs[0]
                p.style = source_style
                p.paragraph_format.line_spacing = 1.05
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.keep_with_next = False
                p.paragraph_format.keep_together = True
                if colidx > 0: p.alignment = 1
                r = p.add_run(lines[source_line-1])
                r.font.size = Pt(11)
                r.bold = rowidx == 0
                source_paragraphs.append(p)
        pr = table._tbl.tblPr
        border = OxmlElement('w:tblBorders')
        for edge in ['top','bottom','insideH']:
            e = OxmlElement('w:'+edge)
            e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'5')
            e.set(qn('w:color'),'777777'); border.append(e)
        pr.append(border)
        add_notes(38)
        continue
    p = doc.add_paragraph(style='Original table line' if 21 <= lineno <= 38 else 'Original source')
    p.paragraph_format.widow_control = True
    if lineno in (3,15):
        p.paragraph_format.keep_with_next = True
        p.add_run(text).bold = True
    elif lineno == 18:
        equation(p)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.keep_together = True
    else:
        p.add_run(text)
    if lineno == 38: p.paragraph_format.keep_with_next = False
    source_paragraphs.append(p)
    add_notes(lineno)

doc.add_heading('Requirements affecting the complete manuscript', 1)
doc.add_paragraph('The pasted material contains Sections 2–3. The absence here of an abstract, introduction, conclusion, full references, or AI statement is not evidence that the full paper omits them. Check those components in the complete manuscript. [E5]')
checks = [
 'Track A aims for 10–12 main-text pages, excluding references and appendices. The guideline says the number of pages does not matter; it is an aim, not a rigid cap. Keep replication around 2–3 pages (20–30%) and the extension around 4–6 pages (60–70%). The review copy’s page count does not measure your submission length.',
 'State one main contribution: treatment-effect heterogeneity in recorded departure by baseline writing status. Keep the destination decomposition secondary and put detailed robustness checks in appendices. The excerpt repeats the replication result and model several times (especially lines 43–47); consolidate repetitions so each paragraph advances that point. This follows the explicit concise-writing rule.',
 'Weave relevant literature into the introduction; the instructions ask for no separate literature-review section. Include only works cited in the reference list and use one author-date convention.',
 'The submission is a PDF named Lastname_Paper.pdf. Page 1 contains title, name, and an abstract under 150 words; no separate cover. Use 12-point text, 1.5 spacing, numbered pages, and numbered sections. Remove “(edited),” the date, and @Shay only if those source-header items have entered the manuscript itself.',
 'Resolve footnote markers ¹, ², ¹⁷, and the trailing 6 against the actual full-document footnotes. An unresolved pasted marker is not proof that the manuscript footnote is absent.',
 'The complete AI Use Statement must describe actual coding, analysis, drafting, and editing assistance. The course encourages AI coding assistance and strongly discourages outsourcing substantive writing. It permits controlled polishing with itemized changes; this review leaves the decisions with you. Keep the conversation transcripts.',
 'Ensure submitted code reproduces every final table and figure, with a README containing author, title, environment, and commands. A shareable-folder link is strongly encouraged. Tables and figures must be publication quality; a coefficient plot is recommended, and the final-paper rules prohibit pie charts in the manuscript.'
]
for item in checks:
    p = doc.add_paragraph(style='Review note')
    p.add_run('(Full-paper check: ').bold = True
    p.add_run(item + ' [E5])')

doc.add_heading('Evidence key', 1)
evidence = [
 ('E1', 'Duflo, Esther, Rema Hanna, and Stephen P. Ryan. 2012. “Incentives Work: Getting Teachers to Come to School.” American Economic Review 102(4): 1241–1278. DOI: 10.1257/aer.102.4.1241. Local original: seminar_paper_drafts/00_Original_Paper_Duflo_Hanna_Ryan_2012.pdf. Printed p. 1241 corresponds to physical PDF page 1. Table 2: printed p. 1249; Table 10: printed p. 1274. All page references above use printed article pages.'),
 ('E2', 'Implemented analysis: scripts/audit_transition_extension.py (joins and recodes, lines 100–136; last-up-to-five rule and outcome construction, lines 143–159; interaction model); scripts/audit_transition_scope.py (fixed calendar and block checks; null-imposed wild-bootstrap-t, lines 200–243). These scripts, rather than an invented model, define what was actually fitted.'),
 ('E3', 'Saved results: outputs/research_direction_audit/transition_table10_reproduction.csv; outputs/scope_search_2026_09_22/transition/transition_scope_model_results.csv; transition_scope_group_effects.csv; drop_wildbootstrap_summary.json; drop2_wildbootstrap_summary.json. Attendance reconstruction: seminar_paper_drafts/source/attendance_recheck.json. The primary model is primary_consensus_no_fe, N = 2,195, 111 schools.'),
 ('E4', 'Measurement evidence: outputs/scope_search_2026_09_22/transition/transition_scope_measurement_summary.csv; transition_open_check_coverage.csv; sentinel_audit.md; transition_incomplete_last5_bounds.csv; outputs/scope_search_2026_09_22/measurement/sentinel_interaction_bounds.json; scripts/audit_sentinel_measurement.py. The 66.4% blank-reason calculation is 7,112/10,707 selected check records.'),
 ('E5', 'Assignment: /Users/shayb/Downloads/Experiments_Final_Paper_Instructions_2026 (2).pdf, pp. 1–4, and /Users/shayb/Downloads/Experiments Seminar Paper Instructions 2026 (3).pdf, Track A. Specific requirements above are attributed to the final-paper instructions. Requirements are distinct from the optional clarity advice.'),
 ('E6', 'Replication data: Duflo, Hanna, and Ryan, “Incentives Work: Getting Teachers to Come to School,” Version 1, AEA/ICPSR replication package E112523V1, distributed 2019. DOI: https://doi.org/10.3886/E112523V1. Verify the final bibliography’s dataset-entry format separately from the 2012 journal article.')
]
for key,text in evidence:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.add_run(key + ' — ').bold = True
    p.add_run(text)

doc.save(TARGET)
after = {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in PROTECTED}
assert before == after, 'A protected source changed'
assert [p.text for i,p in enumerate(source_paragraphs,1) if i != 18] == [t for i,t in enumerate(lines,1) if i != 18]
assert not any(p.text.startswith('(') and not p.text.endswith(')') for p in doc.paragraphs if p.style.name == 'Review note')
with ZipFile(TARGET) as z:
    xml = etree.fromstring(z.read('word/document.xml'))
    ns = {'m':'http://schemas.openxmlformats.org/officeDocument/2006/math', 'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    eqs = xml.xpath('//m:oMath',namespaces=ns)
    assert len(eqs) == 2
    assert not xml.xpath('//w:ins | //w:del',namespaces=ns), 'No tracked changes should alter user prose'
    math_tokens = [''.join(e.xpath('.//m:t/text()',namespaces=ns)) for e in eqs]
    assert math_tokens == ['Yisc = αs + γsc + λc + τisc × Tsc + δisc × (Wic × Tsc) + εics     (1)', 'Yis = α + τ Ts + β Wis + δ (Ts × Wis) + εis']
validation = dict(output=str(TARGET), source=str(SOURCE), source_lines=len(lines),
 annotation_count=len(NOTES), categories=dict(Counter(n['category'] for n in NOTES)),
 protected_source_sha256_before=before, protected_source_sha256_after=after,
 all_source_prose_preserved=True, original_equation_terms_preserved=True,
 native_equations=2, review_notes_are_parenthetical=True)
(OUT/'validation.json').write_text(json.dumps(validation,indent=2)+'\n')
(OUT/'annotations.json').write_text(json.dumps(NOTES,indent=2,ensure_ascii=False)+'\n')
print(json.dumps({k:validation[k] for k in ['output','annotation_count','categories','all_source_prose_preserved','native_equations']}))
